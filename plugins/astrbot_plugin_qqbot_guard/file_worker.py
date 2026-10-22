"""Bounded attachment extraction. Never executes document code or formulas."""
import json
import os
from pathlib import Path
import resource
import sys
import zipfile

TEXT_TYPES={'.txt','.md','.markdown','.csv','.tsv','.log','.json','.jsonl','.yaml','.yml',
    '.xml','.html','.htm','.tex','.py','.js','.ts','.tsx','.jsx','.java','.c','.cpp','.h',
    '.hpp','.cs','.go','.rs','.sh','.sql','.r','.rb','.toml','.ini','.conf','.css','.vue','.svelte','.ipynb'}
SUPPORTED_TYPES=TEXT_TYPES|{'.pdf','.docx','.xlsx'}


class ExtractionError(Exception):
    pass


def check_zip(path):
    with zipfile.ZipFile(path) as archive:
        items=archive.infolist()
        if len(items)>5000 or sum(item.file_size for item in items)>48*1024**2:
            raise ExtractionError('archive_too_large')
        if any(item.file_size>16*1024**2 or (item.file_size>1024**2 and item.file_size>300*max(1,item.compress_size)) for item in items):
            raise ExtractionError('archive_too_large')


def extract(path, name, max_chars=24000, image_dir=None, max_images=2):
    path=Path(path)
    suffix=Path(name).suffix.lower()
    if suffix not in SUPPORTED_TYPES:raise ExtractionError('unsupported_type')
    result={'text':'','truncated':False,'notices':[],'images':[],'units_read':0}
    chunks=[]
    used=0
    def add(text):
        nonlocal used
        text=str(text or '').strip()
        if not text:return True
        room=max_chars-used
        if room<=0:
            result['truncated']=True
            return False
        chunks.append(text[:room])
        used+=min(room,len(text))+1
        if len(text)>room:result['truncated']=True
        return len(text)<=room
    if suffix in TEXT_TYPES:
        raw=path.read_bytes()
        if b'\x00' in raw[:2048] and not raw.startswith((b'\xff\xfe',b'\xfe\xff')):
            raise ExtractionError('binary_file')
        text=None
        for encoding in ('utf-8-sig','utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'gb18030'):
            try:text=raw.decode(encoding);break
            except UnicodeError:pass
        if text is None:raise ExtractionError('unknown_encoding')
        add(text)
        result['units_read']=1
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path,strict=False)
        if reader.is_encrypted and not reader.decrypt(''):raise ExtractionError('encrypted_file')
        count=len(reader.pages)
        result['total_pages']=count
        for index,page in enumerate(reader.pages[:min(count,40)]):
            text=page.extract_text() or ''
            result['units_read']+=1
            if not add('[第 '+str(index+1)+' 页]\n'+text):break
            if len(text.strip())<20:
                saved=False
                # A common scan has exactly one full-page raster. Multiple
                # embedded figures do not reconstruct a page's layout.
                images=page.images if image_dir and len(result['images'])<max_images else []
                if len(images)==1:
                    import io
                    from PIL import Image
                    raw=images[0].data
                    if len(raw)<=10*1024**2:
                        with Image.open(io.BytesIO(raw)) as image:
                            if image.width*image.height<=20_000_000:
                                image=image.convert('RGB')
                                image.thumbnail((1600,2000))
                                output=Path(image_dir)/(path.stem+'-scan-'+str(index+1)+'.jpg')
                                image.save(output,'JPEG',quality=85,optimize=True)
                                output.chmod(0o600)
                                result['images'].append(str(output))
                                result['notices'].append('第 '+str(index+1)+' 页为扫描图，已附给识图模型。')
                                saved=True
                if not saved:
                    result['truncated']=True
                    result['notices'].append('第 '+str(index+1)+' 页未提取到完整文字，需提供该页清晰图片。')
        if result['units_read']<count:result['truncated']=True
        if any(len((reader.pages[i].extract_text() or '').strip())<20 for i in range(min(result['units_read'],2))) and not result['images']:
            result['notices'].append('该 PDF 含扫描页，无法仅靠提取文字完整阅读。')
        if not any((reader.pages[i].extract_text() or '').strip() for i in range(result['units_read'])) and not result['images']:
            raise ExtractionError('empty_document')
        result['notices'].append('PDF 文字提取可能丢失公式或版式；未附出的图表不代表已经看过。')
    elif suffix=='.docx':
        check_zip(path)
        from docx import Document
        document=Document(path)
        # Preserve paragraph/table order, and include visible equation tokens.
        for block in document.element.body:
            if block.tag.endswith('}sectPr'):continue
            text=' '.join(node.text for node in block.iter() if node.tag.endswith('}t') and node.text)
            result['units_read']+=1
            if not add(text):break
        result['notices'].append('Word 提取正文、表格与公式文字；图片、批注、页眉页脚和完整公式版式未解析。')
    elif suffix=='.xlsx':
        check_zip(path)
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book.worksheets[:5]:
                if not add('[工作表：'+sheet.title+']'):break
                for row in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row or 1000,1000),max_col=min(sheet.max_column or 20,20),values_only=True):
                    result['units_read']+=1
                    if not add('\t'.join('' if item is None else str(item) for item in row)):break
                if (sheet.max_row or 0)>1000 or (sheet.max_column or 0)>20:result['truncated']=True
                if result['truncated']:break
            if len(book.worksheets)>5:result['truncated']=True
        finally:book.close()
        result['notices'].append('表格最多 5 张工作表、每表 1000 行×20 列；公式仅作为文字读取，未执行或重算。')
    result['text']='\n'.join(chunks)
    if not result['text'].strip() and not result['images']:raise ExtractionError('empty_document')
    if result['truncated']:result['notices'].append('文件超过本次读取范围，以下内容不是全文；需要时请拆分或指定章节。')
    return result


def main():
    resource.setrlimit(resource.RLIMIT_AS,(384*1024**2,384*1024**2))
    resource.setrlimit(resource.RLIMIT_CPU,(12,12))
    resource.setrlimit(resource.RLIMIT_NOFILE,(64,64))
    try:
        result=extract(sys.argv[1],sys.argv[2],int(sys.argv[3]),sys.argv[4],int(sys.argv[5]))
        status=Path('/proc/self/status')
        highwater=next((line.split()[1] for line in status.read_text().splitlines() if line.startswith('VmHWM:')),None) if status.exists() else None
        result['peak_rss_kib']=int(highwater) if highwater is not None else resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        print(json.dumps({'ok':True,'document':result},ensure_ascii=False))
    except Exception as exc:
        code=str(exc) if isinstance(exc,ExtractionError) else type(exc).__name__
        print(json.dumps({'ok':False,'error':code}))


if __name__=='__main__':main()
