import asyncio
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
from types import ModuleType,SimpleNamespace
import unittest
from unittest.mock import AsyncMock,patch
import zipfile

ROOT=Path(__file__).resolve().parent.parent
package=ModuleType('qqbot_attachment_fixture')
package.__path__=[str(ROOT/'plugins/astrbot_plugin_qqbot_guard')]
sys.modules[package.__name__]=package
spec=importlib.util.spec_from_file_location(package.__name__+'.attachments',Path(package.__path__[0])/'attachments.py')
module=importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
worker=sys.modules[package.__name__+'.file_worker']


def pdf(path,text='PDF evidence ALPHA_7632',scan=False,pages=1):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject,NumberObject
    writer=PdfWriter()
    for _ in range(pages):
        page=writer.add_blank_page(width=300,height=250)
        content=DecodedStreamObject()
        if scan:
            image=DecodedStreamObject()
            image.set_data(bytes((255,0,0))*100*80)
            image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),
                NameObject('/Width'):NumberObject(100),NameObject('/Height'):NumberObject(80),
                NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im0'):writer._add_object(image)})})
            content.set_data(b'q 240 0 0 200 20 20 cm /Im0 Do Q')
        else:
            font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
            page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            content.set_data(('BT /F1 12 Tf 20 100 Td ('+text+') Tj ET').encode('ascii'))
        page[NameObject('/Contents')]=writer._add_object(content)
    writer.write(path)


class WorkerTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def test_utf8_text_and_character_limit(self):
        path=self.root/'readme.md'
        path.write_text('中文证据 ALPHA_55\n'+'内容'*100)
        result=worker.extract(path,path.name,40)
        self.assertIn('ALPHA_55',result['text'])
        self.assertTrue(result['truncated'])
        self.assertLessEqual(len(result['text']),40)

    def test_gb18030_and_utf16(self):
        path=self.root/'notes.txt'
        for encoding in ('gb18030','utf-16'):
            path.write_bytes('中文文件：答案 63。'.encode(encoding))
            self.assertIn('答案 63',worker.extract(path,path.name)['text'])

    def test_source_code_is_read_never_executed(self):
        path=self.root/'code.py'
        path.write_text('raise RuntimeError("THIS MUST NOT RUN")\nprint(7)')
        self.assertIn('THIS MUST NOT RUN',worker.extract(path,path.name)['text'])

    def test_unsupported_and_binary_files_are_rejected(self):
        path=self.root/'payload.exe'
        path.write_bytes(b'\x00payload')
        with self.assertRaises(worker.ExtractionError):worker.extract(path,path.name)
        with self.assertRaises(worker.ExtractionError):worker.extract(path,'payload.txt')

    def test_docx_paragraphs_tables_and_equation_text(self):
        from docx import Document
        from docx.oxml import OxmlElement
        document=Document()
        paragraph=document.add_paragraph('测试计划 ALPHA_72')
        math=OxmlElement('m:oMath')
        token=OxmlElement('m:t');token.text='x²+1';math.append(token)
        paragraph._p.append(math)
        table=document.add_table(rows=1,cols=2)
        table.cell(0,0).text='总数';table.cell(0,1).text='63'
        path=self.root/'plan.docx';document.save(path)
        text=worker.extract(path,path.name)['text']
        for item in ('ALPHA_72','x²+1','总数','63'):self.assertIn(item,text)

    def test_xlsx_values_and_formulas_remain_literal(self):
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='预算'
        sheet.append(['项目','数量']);sheet.append(['演示',19]);sheet.append(['合计','=SUM(B2:B2)'])
        path=self.root/'budget.xlsx';book.save(path)
        result=worker.extract(path,path.name)
        self.assertIn('19',result['text']);self.assertIn('=SUM(B2:B2)',result['text'])
        self.assertIn('未执行',''.join(result['notices']))

    def test_pdf_text_is_recovered(self):
        path=self.root/'text.pdf';pdf(path)
        result=worker.extract(path,path.name)
        self.assertIn('ALPHA_7632',result['text'])
        self.assertEqual(result['total_pages'],1)

    def test_blank_and_encrypted_pdf_are_not_fabricated(self):
        from pypdf import PdfWriter
        path=self.root/'empty.pdf'
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(path)
        with self.assertRaises(worker.ExtractionError):worker.extract(path,path.name)
        writer.encrypt('secret');writer.write(path)
        with self.assertRaisesRegex(worker.ExtractionError,'encrypted_file'):worker.extract(path,path.name)

    def test_pdf_scan_is_attached_for_remote_vision(self):
        path=self.root/'scan.pdf';pdf(path,scan=True)
        result=worker.extract(path,path.name,image_dir=self.root)
        self.assertEqual(len(result['images']),1)
        from PIL import Image
        with Image.open(result['images'][0]) as image:
            color=image.getpixel((50,40))
            self.assertGreater(color[0],250)
            self.assertLess(color[1],5);self.assertLess(color[2],5)

    def test_pdf_page_limit_and_scan_budget_are_explicit(self):
        path=self.root/'long.pdf';pdf(path,pages=41)
        result=worker.extract(path,path.name)
        self.assertEqual(result['units_read'],40)
        self.assertTrue(result['truncated'])
        path=self.root/'scans.pdf';pdf(path,scan=True,pages=3)
        result=worker.extract(path,path.name,image_dir=self.root,max_images=1)
        self.assertEqual(len(result['images']),1)
        self.assertTrue(result['truncated'])
        self.assertIn('第 3 页',''.join(result['notices']))

    def test_zip_bomb_is_rejected_before_parser(self):
        path=self.root/'bomb.docx'
        with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml','0'*(2*1024**2))
        with self.assertRaisesRegex(worker.ExtractionError,'archive_too_large'):worker.extract(path,path.name)


class ReaderTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.reader=module.AttachmentReader(self.root/'uploads')
        self.reader.folder.mkdir()

    def part(self,name='readme.txt',text='文件证据 ALPHA_126'):
        path=self.reader.folder/name;path.write_text(text)
        return SimpleNamespace(name=name,file_=str(path),url='')

    async def test_real_worker_reads_and_protects_downloaded_copy(self):
        part=self.part()
        docs=await self.reader.prepare([(part,'direct')])
        self.assertIn('ALPHA_126',docs[0]['text'])
        self.assertTrue(Path(part.file_).is_file())
        self.assertEqual(Path(part.file_).stat().st_mode&0o777,0o600)
        self.assertGreater(docs[0]['peak_rss_kib'],0)

    async def test_two_files_share_total_character_budget(self):
        first=self.part('first.txt','A'*20000);second=self.part('second.txt','B'*20000)
        docs=await self.reader.prepare([(first,'direct'),(second,'quoted')])
        self.assertEqual(len(docs),2)
        self.assertLessEqual(sum(len(doc['text']) for doc in docs),24000)
        self.assertTrue(all(doc['text'] for doc in docs))

    async def test_official_file_uri_with_spaces_is_read(self):
        part=self.part('文件 notes.txt')
        part.file_=Path(part.file_).as_uri()
        docs=await self.reader.prepare([(part,'direct')])
        self.assertIn('ALPHA_126',docs[0]['text'])

    async def test_reading_an_arbitrary_host_file_is_rejected(self):
        outside=self.root/'outside.txt';outside.write_text('DO NOT READ')
        part=SimpleNamespace(name='outside.txt',file_=str(outside),url='')
        with self.assertRaisesRegex(module.AttachmentError,'invalid_attachment'):await self.reader.prepare([(part,'direct')])
        self.assertEqual(outside.read_text(),'DO NOT READ')

    async def test_size_and_file_count_limits(self):
        part=self.part(text='X'*200)
        self.reader.max_bytes=100
        with self.assertRaisesRegex(module.AttachmentError,'file_too_large'):await self.reader.prepare([(part,'direct')])
        with self.assertRaisesRegex(module.AttachmentError,'too_many_files'):await self.reader.prepare([(part,'direct')]*3)

    async def test_private_and_metadata_urls_are_rejected(self):
        for url in ('http://127.0.0.1/a','http://169.254.169.254/latest/meta-data','http://10.0.0.1/a',
                    'http://[::1]/a','http://localhost/a','file:///etc/passwd','http://user:pass@example.com/a'):
            with self.assertRaises(module.AttachmentError):module.validate_url(url)
        self.assertEqual(module.validate_url('https://example.com/file.pdf'),'https://example.com/file.pdf')

    async def test_dns_resolver_rejects_private_destinations(self):
        resolver=module.PublicResolver()
        try:
            resolver.resolver=SimpleNamespace(resolve=AsyncMock(return_value=[{'host':'127.0.0.1'}]),close=AsyncMock())
            with self.assertRaises(module.AttachmentError):await resolver.resolve('public-looking.example',443)
        finally:await resolver.close()

    async def test_redirect_to_internal_address_is_blocked(self):
        class Response:
            status=302
            headers={'Location':'http://169.254.169.254/private'}
            async def __aenter__(self):return self
            async def __aexit__(self,*args):pass
        calls=[]
        class Session:
            def __init__(self,**kwargs):self.connector=kwargs['connector']
            async def __aenter__(self):return self
            async def __aexit__(self,*args):await self.connector.close()
            def get(self,url,**kwargs):calls.append(url);return Response()
        with patch.object(module.aiohttp,'ClientSession',Session):
            with self.assertRaises(module.AttachmentError):await self.reader._download('https://example.com/file',self.reader.folder/'out')
        self.assertEqual(calls,['https://example.com/file'])

    async def test_streaming_download_size_limit_cleans_partial_file(self):
        class Content:
            async def iter_chunked(self,size):
                yield b'X'*80
                yield b'Y'*80
        class Response:
            status=200;content_length=None;content=Content();headers={}
            async def __aenter__(self):return self
            async def __aexit__(self,*args):pass
        class Session:
            def __init__(self,**kwargs):self.connector=kwargs['connector']
            async def __aenter__(self):return self
            async def __aexit__(self,*args):await self.connector.close()
            def get(self,*args,**kwargs):return Response()
        self.reader.max_bytes=100
        part=SimpleNamespace(name='readme.txt',file_='',url='https://example.com/readme.txt')
        with patch.object(module.aiohttp,'ClientSession',Session):
            with self.assertRaisesRegex(module.AttachmentError,'file_too_large'):await self.reader.prepare([(part,'direct')])
        self.assertEqual(list(self.reader.folder.iterdir()),[])

    async def test_cancelled_parser_is_killed(self):
        started=asyncio.Event()
        process=SimpleNamespace(returncode=None,kill=lambda:setattr(process,'returncode',-9),wait=AsyncMock())
        async def communicate():started.set();await asyncio.Event().wait()
        process.communicate=communicate
        with patch.object(module.asyncio,'create_subprocess_exec',AsyncMock(return_value=process)):
            task=asyncio.create_task(self.reader._parse(Path('x'),'x.txt',100,self.reader.folder,0))
            await started.wait();task.cancel()
            with self.assertRaises(asyncio.CancelledError):await task
        self.assertEqual(process.returncode,-9)
        process.wait.assert_awaited_once()

    async def test_prune_ignores_unowned_directories(self):
        important=self.reader.folder/'existing-service-data';important.mkdir()
        (important/'keep.txt').write_text('KEEP')
        self.reader._prune()
        self.assertEqual((important/'keep.txt').read_text(),'KEEP')


if __name__=='__main__':unittest.main()
